"""Nexo Digital - informes descargables: HTML imprimible (con gráficos) y Excel."""
import base64
import io
from html import escape as E
from pathlib import Path

import core
import evidencia
import finanzas
import marca
import narrativa
import nexo
import preguntas

AQUI = Path(__file__).parent
# Paleta institucional UdeA (Manual de Identidad Institucional): verde 349 C, 7740 C, 361 C, lima 375 C, turquesa 7465 C, 7718 C
NAVY, TEAL, AMBER, ROJO, GRIS = "#026937", "#0E7774", "#D97706", "#B42318", "#5B6B63"
LIMA, VERDE_OSC = "#8DC63F", "#014A27"
COL_NIVEL = {"verde": "#12805C", "amarillo": "#B7791F", "rojo": "#B42318", "info": "#64748B"}
TXT_NIVEL = {"verde": "Bien", "amarillo": "Atención", "rojo": "Alerta", "info": "Informativo"}
PELDANOS = {0: "Sin tecnología", 1: "Gratuito o ya disponible", 2: "Bajo costo", 3: "Automatización", 4: "Servicio externo"}
SERIE_COL = [VERDE_OSC, "#35944B", AMBER, ROJO, "#1D4E89", "#7C3AED"]


def _logo_b64():
    try:
        return base64.b64encode((AQUI / "assets" / "logo_horizontal.png").read_bytes()).decode()
    except Exception:
        return ""


def _m(x):
    return core.cop(x)


def _corto(x):
    """Formato corto para ejes: 74,0 M."""
    a = abs(x)
    s = "-" if x < 0 else ""
    if a >= 1e9:
        return f"{s}{a/1e9:.1f} mil M"
    if a >= 1e6:
        return f"{s}{a/1e6:.1f} M"
    if a >= 1e3:
        return f"{s}{a/1e3:.0f} mil"
    return f"{s}{a:.0f}"


# ------------------------------------------------------------------ Gráficos SVG
def svg_barras_mensual(resumen):
    W, H, L, B, T = 760, 280, 60, 50, 14
    n = len(resumen)
    vmax = float(max(resumen["ingreso"].max(), resumen["gasto"].max())) or 1.0
    pw, ph = W - L - 10, H - B - T
    g = pw / n
    bw = g * 0.36
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" class="chart">']
    for i in range(5):
        v = vmax * i / 4
        y = T + ph - ph * i / 4
        out.append(f'<line x1="{L}" y1="{y:.1f}" x2="{W-10}" y2="{y:.1f}" stroke="#DCE7DF"/>'
                   f'<text x="{L-6}" y="{y+4:.1f}" font-size="10" fill="{GRIS}" text-anchor="end">{_corto(v)}</text>')
    paso = max(1, -(-n // 9))
    for i, (per, fila) in enumerate(resumen.iterrows()):
        x = L + i * g + g * 0.12
        hi, hg = ph * fila["ingreso"] / vmax, ph * fila["gasto"] / vmax
        out.append(f'<rect x="{x:.1f}" y="{T+ph-hi:.1f}" width="{bw:.1f}" height="{hi:.1f}" fill="{LIMA}"/>'
                   f'<rect x="{x+bw+1:.1f}" y="{T+ph-hg:.1f}" width="{bw:.1f}" height="{hg:.1f}" fill="{VERDE_OSC}"/>')
        if i % paso == 0:
            out.append(f'<text x="{x+bw:.1f}" y="{H-28}" font-size="9.5" fill="{GRIS}" text-anchor="middle">{E(str(per))}</text>')
    out.append(f'<rect x="{L}" y="{H-9}" width="9" height="9" fill="{LIMA}"/><text x="{L+13}" y="{H-1}" font-size="10" fill="{GRIS}">Ingresos</text>'
               f'<rect x="{L+75}" y="{H-9}" width="9" height="9" fill="{VERDE_OSC}"/><text x="{L+88}" y="{H-1}" font-size="10" fill="{GRIS}">Gastos</text>')
    out.append("</svg>")
    return "".join(out)


def svg_proyeccion(proyecciones):
    W, H, L, B, T = 760, 280, 66, 36, 14
    series = list(proyecciones.items())
    vals = [v for _, p in series for v in p["saldo"]]
    lo, hi = min(min(vals), 0.0), max(max(vals), 0.0)
    if hi == lo:
        hi = lo + 1
    n = len(series[0][1])
    pw, ph = W - L - 10, H - B - T - 44

    def X(i):
        return L + (pw * i / max(n - 1, 1))

    def Y(v):
        return T + ph - ph * (v - lo) / (hi - lo)
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" class="chart">']
    for i in range(5):
        v = lo + (hi - lo) * i / 4
        out.append(f'<line x1="{L}" y1="{Y(v):.1f}" x2="{W-10}" y2="{Y(v):.1f}" stroke="#DCE7DF"/>'
                   f'<text x="{L-6}" y="{Y(v)+4:.1f}" font-size="10" fill="{GRIS}" text-anchor="end">{_corto(v)}</text>')
    out.append(f'<line x1="{L}" y1="{Y(0):.1f}" x2="{W-10}" y2="{Y(0):.1f}" stroke="{ROJO}" stroke-dasharray="4 3"/>')
    for i, per in enumerate(series[0][1].index):
        ancla = "end" if i == n - 1 and n > 1 else "middle"
        out.append(f'<text x="{X(i):.1f}" y="{T+ph+14}" font-size="9.5" fill="{GRIS}" text-anchor="{ancla}">{E(str(per))}</text>')
    for k, (nombre, p) in enumerate(series):
        col = SERIE_COL[k % len(SERIE_COL)]
        pts = " ".join(f"{X(i):.1f},{Y(v):.1f}" for i, v in enumerate(p["saldo"]))
        out.append(f'<polyline points="{pts}" fill="none" stroke="{col}" stroke-width="{2.6 if k == 0 else 1.7}"/>')
        corto = nombre.split(" (")[0].split(":")[0]
        yy = T + ph + 30 + (k // 3) * 12
        xx = L + (k % 3) * 235
        out.append(f'<rect x="{xx}" y="{yy-8}" width="9" height="9" fill="{col}"/>'
                   f'<text x="{xx+13}" y="{yy}" font-size="9.5" fill="{GRIS}">{E(corto)}</text>')
    out.append("</svg>")
    return "".join(out)


def svg_barras_h(items, color=TEAL, fmt=None, ancho_etq=150, vmax=None):
    """items: [(etiqueta, valor)]; valores pueden ser negativos."""
    if not items:
        return ""
    W, fila = 760, 26
    H = fila * len(items) + 8
    vmax = vmax or max(abs(v) for _, v in items) or 1.0
    pw = W - ancho_etq - 110
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" class="chart">']
    for i, (et, v) in enumerate(items):
        y = 4 + i * fila
        w = pw * abs(v) / vmax
        col = ROJO if v < 0 else color
        txt = fmt(v) if fmt else f"{v:,.1f}"
        out.append(f'<text x="{ancho_etq-8}" y="{y+16}" font-size="11" fill="{NAVY}" text-anchor="end">{E(str(et)[:42])}</text>'
                   f'<rect x="{ancho_etq}" y="{y+4}" width="{max(w, 1):.1f}" height="16" fill="{col}" rx="2"/>'
                   f'<text x="{ancho_etq+w+6:.1f}" y="{y+16}" font-size="11" fill="{GRIS}">{E(txt)}</text>')
    out.append("</svg>")
    return "".join(out)


def svg_estacionalidad(indice):
    items = [(m, float(v)) for m, v in indice.items()]
    W, H, L, B = 760, 170, 40, 26
    pw, ph = W - L - 10, H - B - 14
    lo, hi = 0.8, max(1.2, max(v for _, v in items))
    g = pw / len(items)

    def Y(v):
        return 8 + ph - ph * (v - lo) / (hi - lo)
    out = [f'<svg viewBox="0 0 {W} {H}" xmlns="http://www.w3.org/2000/svg" class="chart">',
           f'<line x1="{L}" y1="{Y(1):.1f}" x2="{W-10}" y2="{Y(1):.1f}" stroke="{GRIS}" stroke-dasharray="4 3"/>',
           f'<text x="{L-6}" y="{Y(1)+4:.1f}" font-size="10" fill="{GRIS}" text-anchor="end">1,0</text>']
    for i, (m, v) in enumerate(items):
        x = L + i * g + g * 0.2
        y = Y(max(v, lo))
        col = TEAL if v >= 1 else AMBER
        out.append(f'<rect x="{x:.1f}" y="{min(y, Y(1)):.1f}" width="{g*0.6:.1f}" height="{abs(Y(1)-y):.1f}" fill="{col}"/>'
                   f'<text x="{x+g*0.3:.1f}" y="{H-8}" font-size="10" fill="{GRIS}" text-anchor="middle">{m}</text>')
    out.append("</svg>")
    return "".join(out)


# ------------------------------------------------------------------ HTML
CSS = f"""
@page {{ size: A4; margin: 16mm 14mm; }}
* {{ box-sizing: border-box; }}
body {{ font-family: 'Segoe UI', Calibri, Arial, sans-serif; color: #14281D; margin: 0; background: #EEF5EF; line-height: 1.5; font-size: 14px; }}
.hoja {{ max-width: 900px; margin: 0 auto; background: #fff; padding: 36px 44px; }}
h1 {{ color: {NAVY}; font-size: 26px; margin: 0 0 4px; }}
h2 {{ color: {NAVY}; font-size: 18px; border-bottom: 2px solid {LIMA}; padding-bottom: 4px; margin: 30px 0 12px; }}
h3 {{ color: {NAVY}; font-size: 14.5px; margin: 16px 0 6px; }}
.portada {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 3px solid {NAVY}; padding-bottom: 14px; margin-bottom: 18px; }}
.portada img {{ height: 50px; }}
.meta {{ color: {GRIS}; font-size: 12.5px; text-align: right; }}
.sem {{ padding: 14px 18px; border-radius: 8px; color: #fff; margin: 10px 0 14px; }}
.sem b {{ font-size: 19px; }}
.kpis {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 10px; margin: 10px 0; }}
.kpi {{ border: 1px solid #C9D8CE; border-radius: 8px; padding: 10px 12px; border-top: 3px solid {LIMA}; }}
.kpi .v {{ font-size: 17px; font-weight: 700; color: {NAVY}; }}
.kpi .t {{ font-size: 11.5px; color: {GRIS}; }}
table {{ width: 100%; border-collapse: collapse; margin: 8px 0 12px; font-size: 12.5px; }}
th {{ background: {NAVY}; color: #fff; text-align: left; padding: 6px 8px; font-weight: 600; }}
td {{ padding: 6px 8px; border-bottom: 1px solid #DCE7DF; vertical-align: top; }}
tr {{ page-break-inside: avoid; }}
.pill {{ display: inline-block; padding: 1px 9px; border-radius: 10px; color: #fff; font-size: 11.5px; font-weight: 600; }}
.nota {{ background: #F4F8F5; border-left: 4px solid {LIMA}; padding: 8px 12px; margin: 8px 0; font-size: 12.5px; color: #334155; }}
.alerta {{ border-left-color: {ROJO}; background: #FEF3F2; }}
.chart {{ width: 100%; height: auto; margin: 6px 0; }}
.pie {{ font-size: 11px; color: {GRIS}; border-top: 1px solid #C9D8CE; margin-top: 30px; padding-top: 8px; }}
.btn {{ position: fixed; top: 14px; right: 14px; background: {TEAL}; color: #fff; border: 0; padding: 10px 16px; border-radius: 6px; font-size: 14px; cursor: pointer; }}
ul {{ margin: 4px 0 8px 20px; padding: 0; }}
h2, h3 {{ page-break-after: avoid; }}
@media print {{ body {{ background: #fff; }} .hoja {{ padding: 0; max-width: none; }} .btn {{ display: none; }}
  .sem, .pill, th, .kpi {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }} }}
@media (max-width: 640px) {{ .kpis {{ grid-template-columns: repeat(2, 1fr); }} .hoja {{ padding: 20px; }} }}
"""


def _tabla(cabeza, filas):
    h = "".join(f"<th>{E(c)}</th>" for c in cabeza)
    b = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in f) + "</tr>" for f in filas)
    return f"<table><tr>{h}</tr>{b}</table>"


def _kpi(valor, titulo):
    return f'<div class="kpi"><div class="v">{E(str(valor))}</div><div class="t">{E(titulo)}</div></div>'


def _pill(nivel):
    return f'<span class="pill" style="background:{COL_NIVEL[nivel]}">{TXT_NIVEL[nivel]}</span>'


def _pres_corto(r):
    if r["cap"] is None:
        return "Sin datos"
    return (_m(r["presupuesto"]) + " / mes") if r["presupuesto"] > 0 else "$0 por ahora"


def _parrafos(lista):
    return "".join(f"<p>{E(t)}</p>" for t in lista)


def _lect(r, titulo):
    for t, x in r["lectura_fin"]:
        if t == titulo:
            return f'<div class="nota"><b>{E(t)}.</b> {E(x)}</div>'
    return ""


def _respaldo_corto(ids):
    refs_ = evidencia.obtener(ids)
    return "; ".join(evidencia.corta(x) for x in refs_)


def _expertos(r):
    """[(recomendación, [refs])] solo para los hallazgos que disparó el caso."""
    out, vistos = [], set()
    for x in r["disparadas"]:
        refs_ = evidencia.obtener(x.get("refs_ids", []))
        if not refs_:
            continue
        titulo = x.get("texto") or x["fuente"]
        if titulo in vistos:
            continue
        vistos.add(titulo)
        out.append((titulo.strip("¿?"), refs_))
    return out


def informe_html(r):
    fin, re_ = r["fin"], r["resumen_ejecutivo"]
    P = []
    logo = _logo_b64()
    P.append(f'<div class="portada"><div>{"<img alt=\"Nexo Digital\" src=\"data:image/png;base64," + logo + "\">" if logo else "<b>NEXO DIGITAL</b>"}'
             f'<h1 style="margin-top:10px">Informe de diagnóstico de transformación digital</h1><div style="font-size:16px;color:{TEAL};font-weight:600">{E(r["caso"])}</div></div>'
             f'<div class="meta">Fecha: {E(r["fecha"])}<br>Sector: {E(r["sector"] or "no indicado")}<br>Confiabilidad del diagnóstico: {E(r["confianza"].split(":")[0])}</div></div>')

    # 1. Resumen ejecutivo
    P.append("<h2>1. Resumen ejecutivo</h2>")
    sem = re_["semaforo"]
    if sem:
        col = [COL_NIVEL["verde"], COL_NIVEL["amarillo"], COL_NIVEL["rojo"]][sem["nivel"]]
        razones = "".join(f"<li>{E(t)}</li>" for _, t in sem["razones"])
        P.append(f'<div class="sem" style="background:{col}"><b>Semáforo financiero: {E(sem["etiqueta"])}</b><ul>{razones}</ul></div>')
    else:
        P.append('<div class="nota">No se calculó el semáforo financiero porque no se cargaron datos financieros.</div>')
    ind, mat = re_["indice"], re_["matriz"]
    P.append(f'<div class="kpis">{_kpi(f"{ind["valor"]:.0f} / 100", "Índice de preparación digital (" + ind["banda"] + ")")}'
             f'{_kpi(mat["nombre"] if mat else "Sin datos", "Estrategia sugerida")}'
             f'{_kpi(_pres_corto(r), "Presupuesto digital prudente")}'
             f'{_kpi(len(r["orden"]), "Frentes con brechas")}</div>')
    P.append("<h3>Conclusión del diagnóstico</h3>" + _parrafos(r["conclusion"]))
    P.append("<h3>Tres prioridades</h3><ol>" + "".join(f"<li><b>{E(n)}.</b> {E(t)}</li>" for n, t in re_["prioridades"]) + "</ol>")
    if re_["primer_paso"]:
        P.append(f'<div class="nota"><b>Primer paso en 30 días:</b> {E(re_["primer_paso"])}</div>')
    if re_["alerta"]:
        P.append(f'<div class="nota alerta"><b>Alerta principal:</b> {E(re_["alerta"])}</div>')

    # 2. Perfil
    if r["perfil_empresa"]:
        P.append("<h2>2. Perfil de la empresa y su contexto</h2>")
        for t, x in r["perfil_empresa"]:
            P.append(f'<p><b>{E(t)}.</b> {E(x)}</p>')

    # 3. Análisis financiero
    P.append("<h2>3. Análisis financiero</h2>")
    mov = fin["mov"]
    if mov:
        if fin["aviso"]:
            P.append(f'<div class="nota">{E(fin["aviso"])}</div>')
        cob = mov["cobertura_meses"]
        P.append('<div class="kpis">'
                 + _kpi(_m(mov["ingreso_prom"]), "Ingresos por mes (promedio)")
                 + _kpi(_m(mov["gasto_prom"]), "Gastos por mes (promedio)")
                 + _kpi(f'{mov["margen"]:.1%}', "Margen de caja (últimos meses)")
                 + _kpi("sin dato" if cob is None else f"{cob:.1f} meses", "Cobertura de caja") + "</div>")
        P.append(_lect(r, "Resultado de caja") + _lect(r, "Colchón de caja") + _lect(r, "Regularidad")
                 + _lect(r, "Estabilidad de los ingresos") + _lect(r, "Tendencia reciente"))
        P.append("<h3>Ingresos y gastos por mes</h3>" + svg_barras_mensual(fin["resumen"]))
        eq = fin["equilibrio"]
        if eq and eq.get("ventas_equilibrio"):
            P.append("<h3>Punto de equilibrio</h3>")
            P.append(_tabla(["Concepto", "Valor mensual"], [
                ["Ingresos promedio", _m(eq["ingreso"])], ["Gastos fijos estimados", _m(eq["fijos"])],
                ["Gastos variables estimados", _m(eq["variables"])],
                ["Margen de contribución", f'{eq["margen_contribucion"]:.1%}'],
                ["<b>Ventas de equilibrio</b>", f'<b>{_m(eq["ventas_equilibrio"])}</b>']]))
            P.append(_lect(r, "Punto de equilibrio"))
            P.append(f'<div class="nota">Los gastos fijos se estiman por el nombre de la categoría '
                     f'({E(", ".join(eq["categorias_fijas"]) or "ninguna detectada")}); es una aproximación, no una medición contable.</div>')
        cat = fin["categorias"]
        if cat is not None and len(cat):
            P.append("<h3>En qué se va el dinero</h3>" + svg_barras_h(
                [(a, b * 100) for a, b in zip(cat["categoria"], cat["participacion"])], NAVY, lambda v: f"{v:.0f} %"))
            P.append(_lect(r, "Estructura de gastos"))
        for a in fin["anomalias"][:3]:
            P.append(f'<div class="nota">Gasto inusual: <b>{E(str(a["categoria"]))}</b> subió {a["pct"]:+.0%} en el último mes frente a su promedio.</div>')
        esc = fin["escenarios"]
        if esc:
            P.append("<h3>Resistencia de la caja: escenarios a 6 meses</h3>" + svg_proyeccion(esc["proyecciones"]))
            t = esc["tabla"]
            P.append(_tabla(["Escenario", "Saldo final", "Saldo mínimo", "Mes en que se agota"],
                            [[E(x.escenario), _m(x.saldo_final), _m(x.saldo_minimo), E(str(x.mes_critico))] for x in t.itertuples()]))
            P.append(_lect(r, "Resistencia de la caja"))
            sens = esc["sensibilidad"]
            P.append("<h3>¿Qué factor pesa más? (efecto en el saldo final)</h3>" + svg_barras_h(
                [(x.factor.split(" (")[0], x.impacto) for x in sens.itertuples()], ROJO, _corto, 190))
        elif fin["error_proy"]:
            P.append(f'<div class="nota">{E(fin["error_proy"])}</div>')
        est = fin["estacionalidad"]
        if est:
            P.append("<h3>Estacionalidad de los ingresos</h3>" + svg_estacionalidad(est["indice"]))
            P.append(_lect(r, "Estacionalidad"))
        if r["palancas"]:
            P.append("<h3>Palancas de mejora cuantificadas</h3>")
            P.append(_tabla(["Palanca", "Efecto mensual", "Efecto anual", "Lectura"],
                            [[E(x["nombre"]), _m(x["mensual"]), _m(x["anual"]), E(x["lectura"])] for x in r["palancas"]]))
            P.append('<div class="nota">Son ejercicios de sensibilidad con las cifras cargadas, no promesas de resultado.</div>')
    else:
        P.append(_parrafos([x for t, x in r["lectura_fin"] if t == "Sin datos financieros"]) or
                 '<div class="nota">No se cargaron movimientos.</div>')

    ests = fin["estados"]
    if ests:
        P.append("<h3>Indicadores de los estados financieros</h3>")
        filas = []
        for it in ests["items"]:
            v = it["valor"]
            vtxt = _m(v) if it["clave"] == "capital_trabajo" else f"{v:,.2f}"
            filas.append([E(it["nombre"]), vtxt, _pill(it["nivel"]), E(it["lectura"]), E(it["definicion"])])
        P.append(_tabla(["Indicador", "Valor", "Lectura", "Interpretación", "Cómo se calcula"], filas))
        P.append(_lect(r, "Estados financieros") + _lect(r, "Origen de la rentabilidad (DuPont)"))
        if r.get("cartera_10"):
            P.append(f'<div class="nota">Cada 10 días menos de cartera libera cerca de <b>{_m(r["cartera_10"])}</b> de caja (ingresos anuales / 365 x 10).</div>')
        P.append('<div class="nota">Los semáforos usan referencias generales de análisis financiero. La comparación más útil es contra el promedio del sector.</div>')
    P.append(_lect(r, "Consistencia entre fuentes"))
    for a in fin["alertas"]:
        P.append(f'<div class="nota{"" if a["nivel"] == "verde" else " alerta"}"><b>{E(a["titulo"])}.</b> {E(a["detalle"])} <i>Acción: {E(a["accion"])}</i></div>')

    # 4. Madurez digital
    P.append("<h2>4. Preparación digital</h2>")
    P.append(f'<p>El índice de preparación digital es de <b>{ind["valor"]:.0f} sobre 100 (nivel {E(ind["banda"].lower())})</b>. '
             'Se calcula con las 30 respuestas del cuestionario (Sí = 2 puntos, En parte = 1, No = 0) y promedia con el mismo peso las ocho secciones. '
             'Las secciones siguen la estructura del Chequeo Digital de MinTIC y el BID; Nexo agrega el contraste con datos financieros reales.</p>')
    P.append(svg_barras_h([(d["seccion"], d["pct"]) for d in r["dimensiones"]], LIMA, lambda v: f"{v:.0f} %", 250, 100.0))
    for d in r["dimensiones"]:
        P.append(f'<h3>{E(d["seccion"])}: {d["pct"]:.0f} % ({E(d["banda"])})</h3><p>{E(d["texto"])}</p>')
        if d["fortalezas"]:
            P.append("<p><b>Fortalezas:</b></p><ul>" + "".join(f"<li>{E(x)}</li>" for x in d["fortalezas"]) + "</ul>")
        if d["brechas"]:
            P.append("<p><b>Brechas:</b></p><ul>" + "".join(f"<li>{E(x)}</li>" for x in d["brechas"]) + "</ul>")
    if mat:
        P.append(f'<div class="nota"><b>Matriz de capacidad financiera y preparación digital.</b> Capacidad financiera {"alta" if mat["cap_alta"] else "limitada"} y '
                 f'preparación digital {"alta" if mat["dig_alta"] else "limitada"}: estrategia <b>{E(mat["nombre"])}</b>. {E(mat["texto"])}</div>')
    P.append("<h3>Contraste de las respuestas con los datos</h3>" + "".join(f'<div class="nota">{E(n)}</div>' for n in r["notas"]))

    # 5. Cuestionario
    P.append("<h2>5. Lectura de cada respuesta</h2>")
    colr = {"Sí": COL_NIVEL["verde"], "En parte": COL_NIVEL["amarillo"], "No": COL_NIVEL["rojo"]}
    filas = [[E(x["seccion"]), E(x["pregunta"]), f'<span class="pill" style="background:{colr[x["respuesta"]]}">{x["respuesta"]}</span>', E(x["lectura"])]
             for x in r["lectura_preguntas"]]
    P.append(_tabla(["Sección", "Pregunta", "Respuesta", "Lectura"], filas))

    # 6. Prioridades
    P.append("<h2>6. Prioridades y recomendaciones</h2>")
    for i, n in enumerate(r["orden"], 1):
        P.append(f"<h3>{i}. {E(nexo.NECESIDADES[n])} (prioridad {r['puntaje'][n]})</h3><ul>")
        for x in [x for x in r["disparadas"] if x["nec"] == n]:
            marco = ' <i>(marco normativo)</i>' if x["tag"] == "normativo" else ""
            P.append(f"<li>{E(x['por_que'])}{marco}" + (f"<br><span style='color:{TEAL}'>{E(x['detalle'])}</span>" if x.get("detalle") else "") + "</li>")
        P.append("</ul>")
    if not r["orden"]:
        P.append("<p>No se detectaron brechas prioritarias con las respuestas dadas.</p>")

    # 7. Escalera de gasto
    P.append("<h2>7. Escalera de gasto: qué hacer y qué esperar</h2>")
    P.append("<p>Para cada frente se muestra primero lo más económico que lo resuelve. Los precios son referenciales: confirme condiciones con el proveedor.</p>")
    for n in r["orden"][:6]:
        filas = []
        for o in r["ops"].get(n, []):
            est = "Recomendado" if o["ok"] else "Aplazar"
            col = COL_NIVEL["verde"] if o["ok"] else COL_NIVEL["rojo"]
            det = E(o["descripcion"])
            if o["motivo"]:
                det += f'<br><i>{E(o["motivo"])}</i>'
            if o["advertencia"]:
                det += f'<br><span style="color:{GRIS}">{E(o["advertencia"])}</span>'
            filas.append([f'<span class="pill" style="background:{col}">{est}</span>', f'{o["peldano"]} · {PELDANOS[o["peldano"]]}',
                          f'<b>{E(o["nombre"])}</b><br>{det}', E(o["costo"])])
        if filas:
            P.append(f"<h3>{E(nexo.NECESIDADES[n])}</h3>" + _tabla(["Estado", "Peldaño", "Opción", "Costo"], filas))
    if r["nct"]:
        P.append("<h3>Inversiones que conviene aplazar</h3>" + _tabla(["Frente", "Opción", "Motivo para aplazar"],
                 [[E(a), E(b), E(c)] for a, b, c in r["nct"]]))

    if r.get("herramientas"):
        P.append("<h3>Herramientas concretas para cada frente</h3><p style=\"color:#5B6B63\">Se ofrecen primero las gratuitas; las de pago solo si hay presupuesto prudente. Confirme condiciones y precios en el sitio oficial.</p>" +
                 _tabla(["Frente", "Herramienta", "Costo", "Para qué sirve"], [[E(nexo.NECESIDADES[x["nec"]]), E(x["nombre"]), E(x["tipo"]), E(x["para"])] for x in r["herramientas"]]))

    # 8. Ruta
    P.append("<h2>8. Ruta de 30 y 90 días</h2><h3>Próximos 30 días</h3><ul>")
    for nec, ops in r["d30"]:
        P.append(f"<li><b>{E(nec)}:</b> " + (E("; ".join(f"{o['nombre']} ({o['costo']})" for o in ops)) or "sin opción disponible en el catálogo") + "</li>")
    P.append("</ul><h3>Días 31 a 90 (solo si lo anterior funcionó y resulta insuficiente)</h3><ul>")
    for nec, op, motivo in r["d90"]:
        t = f"{op['nombre']} ({op['costo']})" if op else (f"aún no. {motivo}" if motivo else "sin paso siguiente en el catálogo")
        P.append(f"<li><b>{E(nec)}:</b> {E(t)}</li>")
    P.append("</ul><h3>Agenda sugerida</h3>")
    for cuando, titulo, acciones in r["agenda"]:
        P.append(f"<p><b>{E(cuando)}.</b> {E(titulo)}</p>")
        if acciones:
            P.append("<ul>" + "".join(f"<li>{E(a)}</li>" for a in acciones) + "</ul>")
    P.append("<h3>Indicadores de seguimiento</h3>" + _tabla(["Frente", "Objetivo", "Indicador", "Meta", "Error frecuente"],
             [[E(x["necesidad"]), E(x["objetivo"]), E(x["indicador"]), E(x["meta"]), E(x["error"])] for x in r["seguimiento"]]))

    # 9. Riesgos y alcance
    P.append("<h2>9. Riesgos de una transformación mal planteada</h2><ul>")
    for t, nota in nexo.RIESGOS:
        P.append(f"<li><b>{E(t)}.</b> {E(nota)}</li>")
    P.append("</ul><h2>10. Cómo se elaboró este diagnóstico</h2><ul>"
             "<li>El índice de preparación digital resume 30 respuestas en ocho secciones; es descriptivo y sirve para comparar la situación de una misma empresa en el tiempo.</li>"
             "<li>Las recomendaciones salen de reglas explícitas que cruzan las respuestas, los datos financieros cargados y el contexto de la empresa.</li>"
             f"<li>Presupuesto prudente: hasta {nexo.PRUDENCIA_EXCEDENTE:.0%} del excedente mensual, y solo si la caja cubre al menos {nexo.MESES_COLCHON_MIN:.0f} mes de gastos.</li>"
             "<li>Los escenarios son ejercicios de sensibilidad, no predicciones. Los gastos fijos se estiman por el nombre de la categoría.</li>"
             "<li>Este informe orienta decisiones y no reemplaza la asesoría de un contador, un abogado o un consultor.</li></ul>")
    ex = _expertos(r)
    if ex:
        P.append("<h2>Anexo. Lo que sostienen los expertos sobre sus recomendaciones</h2>")
        P.append('<p style="color:#5B6B63">Complemento de lectura para quien desee profundizar. Cada texto resume lo que señala la fuente y dónde puede consultarse.</p>')
        for titulo, refs_ in ex:
            P.append(f"<h3>{E(titulo)}</h3>")
            for ref in refs_:
                if ref.get("frase"):
                    P.append(f'<p style="margin:6px 0"><b>{E(evidencia.corta(ref))}</b> señala que {E(ref["frase"])} '
                             f'<span style="color:{GRIS}">Dónde: {E(ref["donde"])}. {E(ref["cita"])} {E(ref["url"])}</span></p>')
                else:
                    P.append(f'<p style="margin:6px 0"><b>Lectura recomendada:</b> {E(ref["cita"])} '
                             f'<span style="color:{GRIS}">{E(ref["url"])}</span></p>')
    P.append(f'<div class="pie"><b>{E(marca.PIE_LINEA1)}</b><br>{E(marca.PIE_LINEA2)}<br>{E(marca.PIE_LINEA3)}</div>')
    cuerpo = "".join(P)
    return (f'<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>Informe Nexo Digital - {E(r["caso"])}</title><style>{CSS}</style></head><body>'
            f'<button class="btn" onclick="window.print()">Imprimir o guardar como PDF</button><div class="hoja">{cuerpo}</div></body></html>')


# ------------------------------------------------------------------ Excel
def informe_excel(r):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    HEAD = PatternFill("solid", fgColor="026937")
    fin, re_ = r["fin"], r["resumen_ejecutivo"]

    def hoja(nombre, cabeza, filas, anchos, nuevo=True):
        ws = wb.create_sheet(nombre) if nuevo else wb.active
        if not nuevo:
            ws.title = nombre
        ws.append(cabeza)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for f in filas:
            ws.append(f)
        for i, a in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = a
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        return ws

    sem = re_["semaforo"]
    filas = [["Empresa", r["caso"]], ["Fecha", r["fecha"]], ["Sector", r["sector"] or "no indicado"],
             ["Semáforo financiero", sem["etiqueta"] if sem else "Sin datos financieros"]]
    if sem:
        filas += [["  Razón", t] for _, t in sem["razones"]]
    filas += [["Índice de preparación digital", f"{r['indice']['valor']:.0f} / 100 ({r['indice']['banda']})"]]
    if r["matriz"]:
        filas += [["Estrategia sugerida", f"{r['matriz']['nombre']}: {r['matriz']['texto']}"]]
    filas += [["Presupuesto digital prudente", re_["presupuesto"]], ["Primer paso en 30 días", re_["primer_paso"]],
              ["Alerta principal", re_["alerta"]], ["Confiabilidad del diagnóstico", r["confianza"]]]
    filas += [["Conclusión", t] for t in r["conclusion"]]
    filas += [[f"Prioridad {i}", f"{n}. {t}"] for i, (n, t) in enumerate(re_["prioridades"], 1)]
    hoja("Resumen", ["Concepto", "Detalle"], filas, [32, 110], nuevo=False)
    if r["perfil_empresa"]:
        hoja("Perfil", ["Aspecto", "Lectura"], [[t, x] for t, x in r["perfil_empresa"]], [34, 120])

    mov = fin["mov"]
    if mov:
        res = fin["resumen"]
        ws = hoja("Mensual", ["Mes", "Ingresos", "Gastos", "Neto"],
                  [[str(p), float(f["ingreso"]), float(f["gasto"]), float(f["neto"])] for p, f in res.iterrows()], [14, 18, 18, 18])
        n = len(res) + 1
        ws.append(["Promedio", f"=AVERAGE(B2:B{n})", f"=AVERAGE(C2:C{n})", f"=AVERAGE(D2:D{n})"])
        for fila in ws.iter_rows(min_row=2, min_col=2):
            for c in fila:
                c.number_format = '#,##0'
        eq = fin["equilibrio"]
        if eq and eq.get("ventas_equilibrio"):
            filas = [["Ingresos promedio mensual", eq["ingreso"]], ["Gastos fijos estimados", eq["fijos"]],
                     ["Gastos variables estimados", eq["variables"]], ["Margen de contribución", eq["margen_contribucion"]],
                     ["Ventas de equilibrio", eq["ventas_equilibrio"]], ["Categorías tomadas como fijas", ", ".join(eq["categorias_fijas"])]]
            hoja("Punto de equilibrio", ["Concepto", "Valor"], filas, [38, 30])
        cat = fin["categorias"]
        if cat is not None and len(cat):
            hoja("Gastos", ["Categoría", "Monto mensual", "Participación", "Acumulado"],
                 [[a, float(b), float(c), float(d)] for a, b, c, d in zip(cat["categoria"], cat["monto_mensual"], cat["participacion"], cat["acumulado"])],
                 [30, 18, 14, 14])
        if fin["escenarios"]:
            t = fin["escenarios"]["tabla"]
            hoja("Escenarios", ["Escenario", "Saldo final", "Saldo mínimo", "Mes en que se agota"],
                 [[x.escenario, float(x.saldo_final), float(x.saldo_minimo), str(x.mes_critico)] for x in t.itertuples()], [60, 18, 18, 22])
            s = fin["escenarios"]["sensibilidad"]
            hoja("Sensibilidad", ["Factor", "Efecto en saldo final"], [[x.factor, float(x.impacto)] for x in s.itertuples()], [60, 24])
    hoja("Lectura financiera", ["Tema", "Lectura"], [[t, x] for t, x in r["lectura_fin"]], [34, 130])
    if r["palancas"]:
        hoja("Palancas", ["Palanca", "Efecto mensual", "Efecto anual", "Lectura"],
             [[x["nombre"], float(x["mensual"]), float(x["anual"]), x["lectura"]] for x in r["palancas"]], [52, 18, 18, 60])
    if fin["estados"]:
        filas = [[it["nombre"], float(it["valor"]), TXT_NIVEL[it["nivel"]], it["lectura"], it["definicion"], it.get("vs_sector", "")]
                 for it in fin["estados"]["items"]]
        hoja("Indicadores", ["Indicador", "Valor", "Lectura", "Interpretación", "Cómo se calcula", "Frente al sector"], filas, [34, 14, 12, 50, 40, 24])
    hoja("Indice digital", ["Sección", "Puntos", "Máximo", "Porcentaje", "Nivel", "Lectura"],
         [[d["seccion"], d["puntos"], d["maximo"], round(d["pct"], 1), d["banda"], d["texto"]] for d in r["dimensiones"]]
         + [["ÍNDICE (promedio simple)", "", "", round(r["indice"]["valor"], 1), r["indice"]["banda"], ""]], [34, 10, 10, 12, 14, 100])
    hoja("Cuestionario", ["Sección", "Pregunta", "Respuesta", "Lectura"],
         [[x["seccion"], x["pregunta"], x["respuesta"], x["lectura"]] for x in r["lectura_preguntas"]], [30, 60, 12, 100])
    filas = [[i, nexo.NECESIDADES[n], x["por_que"], x.get("detalle", ""), "Marco normativo" if x["tag"] == "normativo" else "",
              _respaldo_corto(x.get("refs_ids", []))]
             for i, n in enumerate(r["orden"], 1)
             for x in ([x for x in r["disparadas"] if x["nec"] == n] or [dict(por_que="Por el problema indicado", tag="-", refs_ids=[])])]
    hoja("Prioridades", ["#", "Frente", "Hallazgo", "Qué significa para esta empresa", "Nota", "Lecturas de respaldo"], filas, [5, 34, 70, 70, 16, 40])
    filas = [[nexo.NECESIDADES[n], o["peldano"], PELDANOS[o["peldano"]], o["nombre"], "Recomendado" if o["ok"] else "Aplazar",
              o["costo"], o["motivo"], o["advertencia"], o["url"]]
             for n in r["orden"][:6] for o in r["ops"].get(n, [])]
    hoja("Opciones", ["Frente", "Peldaño", "Nivel", "Opción", "Estado", "Costo", "Motivo", "Aviso", "Enlace"],
         filas, [30, 8, 22, 44, 14, 36, 44, 44, 40])
    filas = [["30 días", nec, "; ".join(f"{o['nombre']} ({o['costo']})" for o in ops) or "sin opción"] for nec, ops in r["d30"]]
    filas += [["31 a 90 días", nec, f"{op['nombre']} ({op['costo']})" if op else (f"aún no. {m}" if m else "sin paso siguiente")]
              for nec, op, m in r["d90"]]
    filas += [[cuando, titulo, "; ".join(acc)] for cuando, titulo, acc in r["agenda"]]
    hoja("Herramientas", ["Frente", "Herramienta", "Costo", "Para qué sirve", "Sitio"],
         [[nexo.NECESIDADES[x["nec"]], x["nombre"], x["tipo"], x["para"], x["url"]] for x in r["herramientas"]], [36, 40, 24, 70, 40])
    hoja("Ruta", ["Horizonte", "Frente", "Acción"], filas, [16, 40, 110])
    hoja("Seguimiento", ["Frente", "Objetivo", "Indicador", "Meta", "Error frecuente"],
         [[x["necesidad"], x["objetivo"], x["indicador"], x["meta"], x["error"]] for x in r["seguimiento"]], [34, 56, 46, 36, 50])
    filas = [[t, evidencia.corta(ref), ref.get("frase") or "Lectura recomendada", ref.get("donde", ""), ref["cita"], ref["url"]] for t, refs_ in _expertos(r) for ref in refs_]
    if filas:
        hoja("Lo que sostienen los expertos", ["Recomendación", "Fuente", "Señala que", "Dónde", "Cita", "Enlace"], filas, [44, 24, 70, 40, 70, 40])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
